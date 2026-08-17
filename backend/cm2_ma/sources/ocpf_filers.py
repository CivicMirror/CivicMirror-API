import io
import zipfile

from openpyxl import load_workbook

from cm2_ingestion.contracts import ContractValidationError
from cm2_ma.constants import OCPF_FILERS_URL
from cm2_ma.source_records import OcpfCandidateRow

from .http import MaPublicBytesSource

REQUIRED_COLUMNS = (
    "CPF ID", "Candidate First Name", "Candidate Last Name",
    "Candidate Street Address", "Candidate City", "Candidate State",
    "Candidate Zip Code", "District Code Sought", "Office Type Sought",
    "District Name Sought", "Closed Date", "Party Affiliation",
)


def _cell_str(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def parse_candidate_filers(zip_bytes: bytes) -> tuple[OcpfCandidateRow, ...]:
    archive = zipfile.ZipFile(io.BytesIO(zip_bytes))
    xlsx_names = [name for name in archive.namelist() if name.endswith(".xlsx")]
    if not xlsx_names:
        raise ContractValidationError("OCPF filers ZIP has no .xlsx entry")

    with archive.open(xlsx_names[0]) as handle:
        workbook = load_workbook(io.BytesIO(handle.read()), read_only=True, data_only=True)

    if "All Candidates" not in workbook.sheetnames:
        raise ContractValidationError("OCPF filer workbook has no 'All Candidates' sheet")
    sheet = workbook["All Candidates"]

    rows_iter = sheet.iter_rows(values_only=True)
    header_row = next(rows_iter, None)
    if header_row is None:
        return ()
    header = [str(cell).strip() if cell is not None else "" for cell in header_row]
    missing = [column for column in REQUIRED_COLUMNS if column not in header]
    if missing:
        raise ContractValidationError("OCPF 'All Candidates' sheet is missing required columns")
    index = {column: header.index(column) for column in header}

    candidates = []
    for row in rows_iter:
        def field(name: str):
            return row[index[name]]

        cpf_id_value = field("CPF ID")
        if cpf_id_value is None:
            continue
        district_code_value = field("District Code Sought")
        district_code_sought = (
            int(district_code_value) if district_code_value not in (None, "") else None
        )

        candidates.append(
            OcpfCandidateRow(
                cpf_id=int(cpf_id_value),
                first_name=_cell_str(field("Candidate First Name")),
                last_name=_cell_str(field("Candidate Last Name")),
                street_address=_cell_str(field("Candidate Street Address")),
                city=_cell_str(field("Candidate City")),
                state=_cell_str(field("Candidate State")),
                zip_code=_cell_str(field("Candidate Zip Code")),
                district_code_sought=district_code_sought,
                office_type_sought=_cell_str(field("Office Type Sought")),
                district_name_sought=_cell_str(field("District Name Sought")),
                closed_date=_cell_str(field("Closed Date")),
                party_affiliation=_cell_str(field("Party Affiliation")),
            )
        )
    return tuple(candidates)


class OcpfFilersSource(MaPublicBytesSource):
    url = OCPF_FILERS_URL

    def parse(self, content: bytes) -> tuple[OcpfCandidateRow, ...]:
        return parse_candidate_filers(content)
