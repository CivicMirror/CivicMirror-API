import io
import zipfile

import pytest
from openpyxl import Workbook

from cm2_ingestion.contracts import ContractValidationError
from cm2_ma.constants import OCPF_FILERS_URL
from cm2_ma.sources.ocpf_filers import OcpfFilersSource, parse_candidate_filers

# Real header, exact column order, captured live from the source workbook's
# "All Candidates" sheet.
_HEADER = (
    "CPF ID", "Account Type Code", "Comm_Name", "Is_Candidate_Only",
    "Candidate First Name", "Candidate Last Name", "Candidate Street Address",
    "Candidate City", "Candidate State", "Candidate Zip Code",
    "Treasurer First Name", "Treasurer Last Name", "Comm Street Address",
    "Comm City", "Comm State", "Comm Zip Code", "Chair First Name",
    "Chair Last Name", "Organization Date", "District Code Sought",
    "Office Type Sought", "District Name Sought", "District Code Held",
    "Office Type Held", "District Name Held", "Closed Date", "Party Affiliation",
)

# Two real sample rows: one open/active (no Closed Date), one closed.
_OPEN_ROW = (
    10010, "U", "Angelo Committee", "False", "Steven", "Angelo", "39 Popmonet Rd",
    "E. Falmouth", "MA", "02536", "Andrew", "Angelo", "60 Halstead Street",
    "Saugus", "MA", "01906", "", "", "2/21/1980", 231, "House", "9th Essex",
    0, "N/A", "No office", "", "Democratic",
)
_CLOSED_ROW = (
    10025, "U", "Berry Committee", "False", "Frederick E.", "Berry",
    "8 Crowninshield Street Unit 410", "Peabody", "MA", "01960",
    "Michael F.", "Tierney", "56 School Street", "Salem", "MA", "01970-2359",
    "", "", "2/5/1982", 107, "Senate", "2nd Essex", 0, "N/A", "No office",
    "7/12/2013", "Democratic",
)


def _filers_zip_bytes(rows: list[tuple]) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_name in ("All Filers", "All Candidates", "PACs", "Local Party Committees"):
        sheet = workbook.create_sheet(sheet_name)
        if sheet_name == "All Candidates":
            sheet.append(_HEADER)
            for row in rows:
                sheet.append(row)
    xlsx_buffer = io.BytesIO()
    workbook.save(xlsx_buffer)

    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, "w") as archive:
        archive.writestr("filer-spreadsheet.xlsx", xlsx_buffer.getvalue())
    return zip_buffer.getvalue()


def test_parses_open_and_closed_rows():
    candidates = parse_candidate_filers(_filers_zip_bytes([_OPEN_ROW, _CLOSED_ROW]))

    assert len(candidates) == 2
    by_cpf_id = {candidate.cpf_id: candidate for candidate in candidates}

    steven = by_cpf_id[10010]
    assert steven.first_name == "Steven"
    assert steven.last_name == "Angelo"
    assert steven.district_code_sought == 231
    assert steven.office_type_sought == "House"
    assert steven.district_name_sought == "9th Essex"
    assert steven.closed_date == ""
    assert steven.party_affiliation == "Democratic"

    frederick = by_cpf_id[10025]
    assert frederick.closed_date == "7/12/2013"
    assert frederick.district_code_sought == 107


def test_rejects_workbook_without_all_candidates_sheet():
    workbook = Workbook()
    xlsx_buffer = io.BytesIO()
    workbook.save(xlsx_buffer)
    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, "w") as archive:
        archive.writestr("filer-spreadsheet.xlsx", xlsx_buffer.getvalue())

    with pytest.raises(ContractValidationError, match="All Candidates"):
        parse_candidate_filers(zip_buffer.getvalue())


def test_rejects_zip_with_no_xlsx_entry():
    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, "w") as archive:
        archive.writestr("readme.txt", "not a workbook")

    with pytest.raises(ContractValidationError, match="no .xlsx entry"):
        parse_candidate_filers(zip_buffer.getvalue())


def test_source_url_and_parse_delegate_correctly():
    source = OcpfFilersSource()
    assert source.url == OCPF_FILERS_URL

    candidates = source.parse(_filers_zip_bytes([_OPEN_ROW]))
    assert len(candidates) == 1
